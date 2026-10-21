"""Ответ модели → сообщения Telegram: Markdown→HTML, разбивка, вынос файлов."""
from __future__ import annotations

import base64
import csv
import html
import io
import re
import secrets

TG_LIMIT = 4000
FILE_PROMPT = (
    "\n\nТы общаешься через Telegram-бота. Бот умеет отправлять пользователю файлы. "
    "Если пользователь просит файл (скрипт, конфиг, документ, таблицу) или содержимое "
    "длиннее ~150 строк, оформи его так:\n"
    '<file name="имя.расширение">\nсодержимое\n</file>\n'
    "Бот превратит блок в настоящий файл. Для .docx пиши содержимое в Markdown, для .xlsx — CSV. "
    "Короткие фрагменты кода оставляй обычными блоками ```."
)

_FILE_RX = re.compile(
    r'<file\s+name=["\']([^"\'<>]+)["\'](?:\s+encoding=["\'](base64)["\'])?\s*>\n?(.*?)\n?</file>',
    re.S | re.I,
)


# ── файлы из ответа ───────────────────────────────────────────
def _strip_fence(s: str) -> str:
    m = re.match(r"^\s*```[^\n]*\n(.*)\n```\s*$", s, re.S)
    return m.group(1) if m else s


def _to_docx(md: str) -> bytes:
    import docx
    doc = docx.Document()
    for line in md.splitlines():
        h = re.match(r"^(#{1,6})\s+(.*)", line)
        if h:
            doc.add_heading(h.group(2), level=min(len(h.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            p = doc.add_paragraph()
            for chunk in re.split(r"(\*\*[^*]+\*\*)", line):
                if chunk.startswith("**") and chunk.endswith("**"):
                    p.add_run(chunk[2:-2]).bold = True
                elif chunk:
                    p.add_run(chunk)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _to_xlsx(text: str) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    for row in csv.reader(io.StringIO(text), dialect):
        ws.append([_num(v) for v in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v: str):
    try:
        return int(v) if re.fullmatch(r"-?\d{1,15}", v) else float(v) if re.fullmatch(r"-?\d+\.\d+", v) else v
    except ValueError:
        return v


def extract_files(text: str) -> tuple[str, list[tuple[str, bytes]]]:
    files: list[tuple[str, bytes]] = []

    def repl(m: re.Match) -> str:
        name, enc, body = m.group(1).strip(), m.group(2), m.group(3)
        name = re.sub(r"[\\/:*?\"<>|]", "_", name) or "file.txt"
        try:
            if enc:
                data = base64.b64decode(body.strip())
            else:
                body = _strip_fence(body)
                low = name.lower()
                if low.endswith(".docx"):
                    data = _to_docx(body)
                elif low.endswith(".xlsx"):
                    data = _to_xlsx(body)
                else:
                    data = body.encode("utf-8")
        except Exception:
            data = body.encode("utf-8")
            name += ".txt"
        files.append((name, data))
        return f"📎 <{name}>"

    return _FILE_RX.sub(repl, text), files


# ── Markdown → Telegram HTML ─────────────────────────────────
def _inline(text: str) -> str:
    stash: dict[str, str] = {}

    def keep(value: str) -> str:
        k = f"\x00{secrets.token_hex(4)}\x00"
        stash[k] = value
        return k

    text = re.sub(r"`([^`\n]+)`", lambda m: keep(f"<code>{html.escape(m.group(1))}</code>"), text)
    text = re.sub(
        r"\[([^\]\n]+)\]\((https?://[^)\s]+)\)",
        lambda m: keep(f'<a href="{html.escape(m.group(2), quote=True)}">{html.escape(m.group(1))}</a>'),
        text,
    )
    text = html.escape(text, quote=False)
    text = re.sub(r"^(#{1,6})\s+(.+)$", r"<b>\2</b>", text, flags=re.M)
    text = re.sub(r"^(\s*)[-*]\s+", r"\1• ", text, flags=re.M)
    text = re.sub(r"^&gt;\s?(.*)$", r"<blockquote>\1</blockquote>", text, flags=re.M)
    text = text.replace("</blockquote>\n<blockquote>", "\n")
    text = re.sub(r"\*\*(?=\S)(.+?)(?<=\S)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"__(?=\S)(.+?)(?<=\S)__", r"<u>\1</u>", text)
    text = re.sub(r"(?<![\w*])\*(?=\S)([^*\n]+?)(?<=\S)\*(?![\w*])", r"<i>\1</i>", text)
    text = re.sub(r"(?<!\w)_(?=\S)([^_\n]+?)(?<=\S)_(?!\w)", r"<i>\1</i>", text)
    text = re.sub(r"~~(?=\S)(.+?)(?<=\S)~~", r"<s>\1</s>", text)
    text = re.sub(r"\|\|(?=\S)(.+?)(?<=\S)\|\|", r"<tg-spoiler>\1</tg-spoiler>", text)
    for k, v in stash.items():
        text = text.replace(k, v)
    return text


def md_to_html(md: str) -> str:
    out = []
    pos = 0
    for m in re.finditer(r"```([\w+#.-]*)[^\n]*\n(.*?)(?:\n```|\Z)", md, re.S):
        out.append(_inline(md[pos:m.start()]))
        lang, code = m.group(1), m.group(2)
        cls = f' class="language-{html.escape(lang)}"' if lang else ""
        out.append(f"<pre><code{cls}>{html.escape(code)}</code></pre>")
        pos = m.end()
    out.append(_inline(md[pos:]))
    return "".join(out).strip()


def split_markdown(md: str, limit: int = TG_LIMIT) -> list[str]:
    """Режет Markdown на куски, не ломая блоки кода (закрывает и переоткрывает ```)."""
    chunks, cur, cur_len = [], [], 0
    fence: str | None = None

    def flush():
        nonlocal cur, cur_len
        if cur:
            body = "\n".join(cur)
            if fence:
                body += "\n```"
            chunks.append(body)
        cur, cur_len = ([fence] if fence else []), (len(fence) + 1 if fence else 0)

    for line in md.split("\n"):
        while len(line) > limit:  # сверхдлинная строка
            if cur_len:
                flush()
            cur.append(line[:limit])
            cur_len += limit
            line = line[limit:]
            flush()
        if cur_len + len(line) + 1 > limit and cur_len:
            flush()
        cur.append(line)
        cur_len += len(line) + 1
        if line.strip().startswith("```"):
            fence = None if fence else line.strip()
    if cur and any(x.strip() for x in cur):
        chunks.append("\n".join(cur))
    return [c for c in chunks if c.strip()]
