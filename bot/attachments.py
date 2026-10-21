"""Разбор входящих файлов: картинки, PDF, DOCX, XLSX и любые текстовые."""
from __future__ import annotations

import csv
import io
import logging
import mimetypes
from pathlib import Path

log = logging.getLogger(__name__)

IMAGE_MIMES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
TEXT_EXT = {
    ".txt", ".md", ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".yaml", ".yml", ".toml", ".ini",
    ".cfg", ".conf", ".csv", ".tsv", ".log", ".xml", ".html", ".htm", ".css", ".scss", ".sql",
    ".sh", ".bash", ".ps1", ".psm1", ".bat", ".cmd", ".c", ".h", ".cpp", ".hpp", ".cs", ".java",
    ".kt", ".go", ".rs", ".rb", ".php", ".swift", ".lua", ".pl", ".r", ".tex", ".env", ".properties",
    ".dockerfile", ".gitignore", ".vue", ".svelte", ".srt", ".vtt", ".rst", ".adoc", ".tf", ".hcl",
}


def guess_mime(name: str, fallback: str | None = None) -> str:
    return fallback or mimetypes.guess_type(name)[0] or "application/octet-stream"


def _decode_text(data: bytes) -> str | None:
    if b"\x00" in data[:4096]:
        return None
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"--- стр. {i} ---\n{page.extract_text() or ''}")
    return "\n".join(pages)


def _docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs]
    for t_idx, table in enumerate(doc.tables, 1):
        out.append(f"\n[Таблица {t_idx}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx(data: bytes, max_rows: int = 2000) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        buf = io.StringIO()
        w = csv.writer(buf)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                buf.write(f"... (обрезано на {max_rows} строках)\n")
                break
            w.writerow(["" if v is None else v for v in row])
        out.append(f"[Лист: {ws.title}]\n{buf.getvalue()}")
    return "\n".join(out)


def extract_text(name: str, mime: str, data: bytes) -> str | None:
    """Текст из документа или None, если формат не поддерживается."""
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf" or mime == "application/pdf":
            return _pdf(data)
        if ext == ".docx":
            return _docx(data)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(data)
        if ext in TEXT_EXT or mime.startswith("text/") or mime in ("application/json", "application/xml"):
            return _decode_text(data)
        # неизвестное расширение — пробуем как текст
        return _decode_text(data)
    except Exception as e:  # битый файл не должен ронять обработку
        log.warning("Не удалось разобрать %s: %s", name, e)
        return None


def is_image(mime: str) -> bool:
    return mime in IMAGE_MIMES
